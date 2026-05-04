from docx import Document

from .base import BaseTool


class DOCXTool(BaseTool):
    """Extrai texto e tabelas de arquivos DOCX."""

    def read(self, file_path: str) -> str:
        path = self._validate_path(file_path)
        doc = Document(str(path))

        parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                parts.append(paragraph.text)

        for table in doc.tables:
            rows = []
            for row in table.rows:
                rows.append(" | ".join(cell.text.strip() for cell in row.cells))
            parts.append("\n".join(rows))

        return "\n\n".join(parts)
